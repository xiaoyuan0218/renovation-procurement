"""商品链接：字段往返、Excel 往返，以及"老文件/老客户端不带这一列时不清空"。

链接是 1.2.5 新加的字段。与定金那次同样的坑：整份覆盖的路径（同步落地、表格
导入）在对方"没带这个字段"时，不能当成"要清空" —— 否则老客户端推一次、老表格
导一份，别人填好的链接就没了。
"""

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import auth
from app.db import SessionLocal, engine
from app.main import app
from app.models import (Allocation, Category, ExtraExpense, Item, ItemList,
                        PurchaseRecord, RecordRoom, Room, User)
from app.seed import init_db
from app.services import excel_io, list_transfer, local_apply
from app.services.link_text import extract_link
from tests.conftest import TEST_PASSWORD as PASSWORD, TEST_USER as USER

LINK = "https://item.jd.com/100012043978.html"

# ---- 用户提供的四条分享文案样例（2026-10-25 记录）----
# 链接本体逐字取自用户提供的四条真实分享文案（京东 3.cn 短链、小红书
# xhslink、淘宝 e.tb.cn、拼多多 mobile.yangkeduo.com，四域互不相干、白名单
# 必漏，故统一走「提取第一个 http(s) URL」）。文案的其余套话部分为等价构造
# ——原文全文未在仓库留存，断言只盯提取结果；将来样例改版以这里为对照基线。
SAMPLES = [
    # 京东：@ 出现在 query 里要完整保留（旧的前缀判断对它无碍，但提取不能截断）
    ("【京东】超值好价 https://3.cn/38K-9pPI?jkl=@YC8fI3EmHbBY@ 复制这条信息后打开京东APP",
     "https://3.cn/38K-9pPI?jkl=@YC8fI3EmHbBY@"),
    # 小红书：链接夹在中文套话中间
    ("看看这个 https://xhslink.com/m/2x3IQ9yoNdZ49IKkEhFMOG 超好用的收纳神器",
     "https://xhslink.com/m/2x3IQ9yoNdZ49IKkEhFMOG"),
    # 淘宝：tk= 参数必须保留，丢了短链就打不开
    ("【淘宝】长按复制整条信息 https://e.tb.cn/h.Tv0QkP?tk=GTkNcXYPqzC 再打开淘宝",
     "https://e.tb.cn/h.Tv0QkP?tk=GTkNcXYPqzC"),
    # 拼多多：整段就是一条纯 URL，应原样入库
    ("https://mobile.yangkeduo.com/goods.html?goods_id=888882334951",
     "https://mobile.yangkeduo.com/goods.html?goods_id=888882334951"),
]


@pytest.fixture()
def client():
    init_db()
    session = SessionLocal()
    for table in (ExtraExpense, RecordRoom, Allocation, Item, PurchaseRecord,
                  Room, Category, User, ItemList):
        session.query(table).delete()
    session.commit()
    session.add(ItemList(name="采购清单", sort=0))
    session.commit()
    session.close()
    auth._failures.clear()
    with TestClient(app) as c:
        c.post("/api/auth/setup", json={"username": USER, "password": PASSWORD})
        yield c
    engine.dispose()


def _hdr(client):
    token = client.post("/api/auth/login",
                        json={"username": USER, "password": PASSWORD}).json()["token"]
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture()
def db():
    init_db()
    session = SessionLocal()
    for table in (RecordRoom, Allocation, PurchaseRecord, Item, Room, Category):
        session.query(table).delete()
    session.commit()
    yield session
    session.close()
    engine.dispose()


@pytest.fixture()
def item(db):
    lst = db.query(ItemList).order_by(ItemList.sort, ItemList.id).first()
    row = Item(name="壁挂炉", unit="台", qty_total=1, price=6999,
               discount_price=5099, list_id=lst.id)
    db.add(row)
    db.commit()
    return row


def test_item_link_roundtrip(client):
    """建物料带链接、读回来还在；换一个再存也跟着变。"""
    h = _hdr(client)
    created = client.post("/api/items", json={
        "name": "小米电视", "unit": "台", "qty_total": 1, "price": 2999,
        "link": LINK, "allocations": [],
    }, headers=h).json()
    assert created["link"] == LINK

    fetched = client.get(f"/api/items/{created['id']}", headers=h).json()
    assert fetched["link"] == LINK

    saved = client.put(f"/api/items/{created['id']}", json={
        "name": "小米电视", "unit": "台", "qty_total": 1, "price": 2999,
        "link": "https://item.taobao.com/item.htm?id=123", "allocations": [],
        "base_rev": fetched["rev"],
    }, headers=h).json()
    assert saved["link"] == "https://item.taobao.com/item.htm?id=123"


def test_link_only_accepts_http_and_https(client):
    """javascript: 之类的伪协议一律清空 —— 它最终会被渲染成可点链接。"""
    h = _hdr(client)
    created = client.post("/api/items", json={
        "name": "危险链接", "unit": "台", "qty_total": 1, "price": 1,
        "link": "javascript:alert(1)", "allocations": [],
    }, headers=h).json()
    assert created["link"] == ""
    assert client.get(f"/api/items/{created['id']}", headers=h).json()["link"] == ""


def test_patch_link_also_only_accepts_http_and_https(client):
    """行内改动（PATCH）走同一条伪协议收口：它同样会存库、被渲染成可点链接。"""
    h = _hdr(client)
    created = client.post("/api/items", json={
        "name": "行内改", "unit": "台", "qty_total": 1, "price": 1, "allocations": [],
    }, headers=h).json()

    cleaned = client.patch(f"/api/items/{created['id']}",
                           json={"link": "javascript:alert(1)"}, headers=h)
    assert cleaned.status_code == 200, cleaned.text
    assert cleaned.json()["link"] == ""

    ok = client.patch(f"/api/items/{created['id']}", json={"link": LINK}, headers=h)
    assert ok.json()["link"] == LINK


def test_excel_roundtrip_keeps_link(db, item):
    """导出带「商品链接」列，紧接着导回来还在。"""
    item.link = LINK
    db.commit()
    data = excel_io.export_xlsx(db, item.list_id)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["物料汇总"]
    header = [c.value for c in ws[1]]
    assert "商品链接" in header
    assert ws.cell(2, header.index("商品链接") + 1).value == LINK

    excel_io.import_template(db, data, mode="replace", list_id=item.list_id)
    db.expire_all()
    assert db.query(Item).one().link == LINK


def test_old_excel_without_link_column_keeps_existing(db, item):
    """老表格（没有「商品链接」列）导回来时，不清空已填的链接。"""
    item.link = LINK
    db.commit()
    data = excel_io.export_xlsx(db, item.list_id)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["物料汇总"]
    header = [c.value for c in ws[1]]
    ws.delete_cols(header.index("商品链接") + 1)
    buf = io.BytesIO()
    wb.save(buf)

    report = excel_io.import_template(db, buf.getvalue(), mode="replace",
                                      list_id=item.list_id)
    assert report.get("warnings") == [], report.get("warnings")
    db.expire_all()
    assert db.query(Item).one().link == LINK


def test_old_client_push_keeps_link(db, item):
    """老客户端推的 payload 里没有 link 字段：按物料名保留已填的链接。"""
    item.link = LINK
    db.commit()
    lst = db.query(ItemList).filter(ItemList.id == item.list_id).one()
    payload = list_transfer.export_list(db, lst)
    for row in payload["items"]:
        row.pop("link", None)

    list_transfer.import_list(db, payload, target=lst)

    db.expire_all()
    assert db.query(Item).one().link == LINK


def test_new_client_push_can_clear_link(db, item):
    """新客户端明确带空串时照清 —— 兼容层只兜"没带"，不拦用户主动清空。"""
    item.link = LINK
    db.commit()
    lst = db.query(ItemList).filter(ItemList.id == item.list_id).one()
    payload = list_transfer.export_list(db, lst)
    for row in payload["items"]:
        row["link"] = ""

    list_transfer.import_list(db, payload, target=lst)

    db.expire_all()
    assert db.query(Item).one().link == ""


def test_samples_extract_from_share_text(client):
    """四条真实分享文案：提取出第一个 http(s) URL 入库，而不是存整段套话。"""
    h = _hdr(client)
    for i, (text, wanted) in enumerate(SAMPLES):
        created = client.post("/api/items", json={
            "name": f"样例{i}", "unit": "个", "qty_total": 1, "price": 1,
            "link": text, "allocations": [],
        }, headers=h)
        assert created.status_code == 200, created.text
        assert created.json()["link"] == wanted, text
        assert client.get(f"/api/items/{created.json()['id']}",
                          headers=h).json()["link"] == wanted


def test_extract_takes_first_of_two_links():
    """文案里有多个链接时取第一个：确定性规则，三端对同一文案必然同结果。"""
    text = ("先看这个 https://xhslink.com/m/2x3IQ9yoNdZ49IKkEhFMOG "
            "再看 https://item.jd.com/100012043978.html 都行")
    assert extract_link(text) == "https://xhslink.com/m/2x3IQ9yoNdZ49IKkEhFMOG"


def test_trailing_brackets_pair_aware():
    """尾部剥括号要配对感知：包裹符摘掉，URL 自带的括号原样保留。"""
    # 外层的「))」是包裹符：剥净
    assert extract_link("see https://x.com/a_(page)) thanks") == "https://x.com/a_(page)"
    assert extract_link("(see https://x.com/a))") == "https://x.com/a"
    # URL 自带的配对括号：一个都不能剥，朴素 rstrip 会把它剥坏
    assert extract_link("https://x.com/a_(page)") == "https://x.com/a_(page)"


def test_extract_idempotent_and_scheme_case():
    """提取对已清洗值幂等（pydantic 入口与落地层双重收口不冲突）；scheme 大小写保留。"""
    for text, wanted in SAMPLES:
        assert extract_link(extract_link(text)) == extract_link(text)
    assert extract_link("HTTP://X.COM/A") == "HTTP://X.COM/A"
    assert extract_link(None) == ""
    assert extract_link("javascript:alert(1)") == ""
    # 裸网址没有 scheme：清空
    assert extract_link("item.jd.com/100012043978.html") == ""


def test_empty_host_values_cleared(client):
    """host 段没有字母数字的值全部清空；形似 host 的正常值照放行。"""
    h = _hdr(client)
    for bad in ("https://", "https://...", "https://)", "https://./x"):
        created = client.post("/api/items", json={
            "name": "空host", "unit": "个", "qty_total": 1, "price": 1,
            "link": bad, "allocations": [],
        }, headers=h).json()
        assert created["link"] == "", bad
    for good in ("https://a", "https://1.2.3.4/x", "https://localhost:8000/x"):
        created = client.post("/api/items", json={
            "name": "好host", "unit": "个", "qty_total": 1, "price": 1,
            "link": good, "allocations": [],
        }, headers=h).json()
        assert created["link"] == good, good


def test_oversize_link_cleared_and_long_text_ok(client):
    """600 字符纯 URL 清空（URL 不能截断）；600 字符文案不再 422，提取出短链。"""
    h = _hdr(client)
    long_url = "https://a" + "x" * 591
    assert len(long_url) == 600
    created = client.post("/api/items", json={
        "name": "超长链接", "unit": "个", "qty_total": 1, "price": 1,
        "link": long_url, "allocations": [],
    }, headers=h).json()
    assert created["link"] == ""

    long_text = "看看这个 " + SAMPLES[0][0] + " " + "好" * 550
    assert len(long_text) > 500
    created = client.post("/api/items", json={
        "name": "超长文案", "unit": "个", "qty_total": 1, "price": 1,
        "link": long_text, "allocations": [],
    }, headers=h)
    assert created.status_code == 200, created.text
    assert created.json()["link"] == SAMPLES[0][1]


def test_sync_push_extracts_link(client):
    """清单级同步入口（POST/PUT /api/sync/lists）推整段文案，落地为提取结果。

    提取发生在 pydantic 校验（import_list 之前的 model_dump），所以必须走
    HTTP 端点测 —— 直喂 dict 调 import_list 不经 pydantic，测不到这条收口。
    """
    h = _hdr(client)
    text = "【京东】超值好价 https://3.cn/38K-9pPI?jkl=@YC8fI3EmHbBY@ 打开京东APP"
    pushed = client.post("/api/sync/lists", json={
        "version": 1,
        "list": {"name": "同步样例", "note": "", "sort": 0, "code": ""},
        "rooms": [], "categories": [],
        "items": [{"name": "样例物料", "unit": "个", "qty_total": 1,
                   "price": 1, "link": text}],
        "expenses": [],
    }, headers=h)
    assert pushed.status_code == 200, pushed.text
    body = pushed.json()
    assert body["payload"]["items"][0]["link"] == SAMPLES[0][1]

    # 覆盖路径（SyncPushIn，手机"以我为准"整份替换）同口径
    replaced = client.put(f"/api/sync/lists/{body['list_id']}", json={
        "version": 1,
        "list": {"name": "同步样例", "note": "", "sort": 0, "code": ""},
        "rooms": [], "categories": [],
        "items": [{"name": "样例物料", "unit": "个", "qty_total": 1,
                   "price": 1, "link": SAMPLES[2][0]}],
        "expenses": [],
        "base_fingerprint": body["fingerprint"],
    }, headers=h)
    assert replaced.status_code == 200, replaced.text
    assert replaced.json()["payload"]["items"][0]["link"] == SAMPLES[2][1]


def test_apply_payload_extracts_link_and_baseline_stays_in_sync(db):
    """桌面落地：整段文案提取入库，且传入的 payload 被就地更新。

    就地更新不是可有可无：_save_binding 把这个 dict 存成三方合并的基线，
    基线与落库不同源时，下一轮合并会把 link 差异误判成"本地改了链接"。
    """
    lst = ItemList(name="提取落地", sort=99)
    db.add(lst)
    db.commit()
    payload = {"list": {}, "items": [
        {"name": "文案", "unit": "个", "qty_total": 1, "price": 1,
         "link": "看看这个 https://xhslink.com/m/2x3IQ9yoNdZ49IKkEhFMOG 收纳",
         "allocations": [], "records": []},
        {"name": "纯URL", "unit": "个", "qty_total": 1, "price": 1,
         "link": "https://item.jd.com/100012043978.html",
         "allocations": [], "records": []},
        {"name": "没带", "unit": "个", "qty_total": 1, "price": 1,
         "allocations": [], "records": []},
    ]}
    local_apply.apply_payload(db, lst, payload)
    db.commit()
    db.expire_all()

    rows = {i.name: i.link
            for i in db.query(Item).filter(Item.list_id == lst.id)}
    assert rows["文案"] == "https://xhslink.com/m/2x3IQ9yoNdZ49IKkEhFMOG"
    assert rows["纯URL"] == "https://item.jd.com/100012043978.html"
    assert rows["没带"] == ""
    # 基线同源：payload 本身也变成了提取后的值
    assert payload["items"][0]["link"] == "https://xhslink.com/m/2x3IQ9yoNdZ49IKkEhFMOG"
    assert payload["items"][1]["link"] == "https://item.jd.com/100012043978.html"


def test_excel_import_extracts_link_and_reports_warning(db, item):
    """表格导入的「商品链接」列是整段文案时：落提取值，并在导入报告里报一声。"""
    item.link = LINK
    db.commit()
    data = excel_io.export_xlsx(db, item.list_id)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["物料汇总"]
    header = [c.value for c in ws[1]]
    link_col = header.index("商品链接") + 1
    ws.cell(2, link_col).value = (
        "看看这个 https://xhslink.com/m/2x3IQ9yoNdZ49IKkEhFMOG 收纳神器")
    buf = io.BytesIO()
    wb.save(buf)

    report = excel_io.import_template(db, buf.getvalue(), mode="replace",
                                      list_id=item.list_id)
    db.expire_all()
    assert db.query(Item).one().link == "https://xhslink.com/m/2x3IQ9yoNdZ49IKkEhFMOG"
    assert any("商品链接" in w for w in report["warnings"]), report["warnings"]

    # 提取不出链接（口令码之类）：清空也要报，不能静默
    item.link = LINK
    db.commit()
    data = excel_io.export_xlsx(db, item.list_id)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["物料汇总"]
    ws.cell(2, link_col).value = "口令：复制这条信息打开淘宝"
    buf = io.BytesIO()
    wb.save(buf)
    report = excel_io.import_template(db, buf.getvalue(), mode="replace",
                                      list_id=item.list_id)
    db.expire_all()
    assert db.query(Item).one().link == ""
    assert any("商品链接" in w for w in report["warnings"]), report["warnings"]
