"""定金与尾款：钱先付、货没到，和普通采购不是一回事。

这条链上每一步都会影响用户看到的钱数，所以逐条盯：
  - 定金只算钱、不推进「已到货数量」，状态不会变成买完；
  - 定金从两个口径的「未付」里都扣掉（付了 1000 定金，未付就该少 1000）；
  - 尾款补齐后回到正常终局：已买完、未付 0、优惠列照旧；
  - 导出导入、回退都把「定金」这个标记一起搬。
"""

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import auth
from app.db import SessionLocal, engine
from app.main import app
from app.models import (Allocation, Category, ExtraExpense, Item, ItemList,
                        PurchaseRecord, RecordRoom, Room, User)
from app.seed import init_db
from app.services import codes, compute, excel_io
from tests.conftest import TEST_PASSWORD as PASSWORD, TEST_USER as USER


@pytest.fixture()
def client():
    init_db()
    session = SessionLocal()
    for table in (ExtraExpense, RecordRoom, Allocation, Item, PurchaseRecord,
                  Room, Category, User, ItemList):
        session.query(table).delete()
    session.commit()
    session.add(ItemList(name="采购清单", sort=0))
    session.commit()
    session.close()
    auth._failures.clear()
    with TestClient(app) as c:
        c.post("/api/auth/setup", json={"username": USER, "password": PASSWORD})
        yield c
    engine.dispose()


def _hdr(client):
    token = client.post("/api/auth/login",
                        json={"username": USER, "password": PASSWORD}).json()["token"]
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture()
def db():
    init_db()
    session = SessionLocal()
    for table in (RecordRoom, Allocation, PurchaseRecord, Item, Room, Category):
        session.query(table).delete()
    session.commit()
    yield session
    session.close()
    engine.dispose()


@pytest.fixture()
def item(db):
    """一件 1 台、原价 6999、日常价 5099 的物料（就是用户报的那个场景）。"""
    lst = db.query(ItemList).order_by(ItemList.sort, ItemList.id).first()
    row = Item(name="壁挂炉", unit="台", qty_total=1, price=6999,
               discount_price=5099, list_id=lst.id)
    db.add(row)
    db.commit()
    return row


def _add(db, item, qty, amount, deposit=False):
    item.records.append(PurchaseRecord(qty=qty, amount=amount, is_deposit=deposit))
    db.commit()
    db.refresh(item)


def _reload(db, item_id):
    db.expire_all()
    return db.get(Item, item_id)


def test_deposit_does_not_mark_as_bought(db, item):
    """只付定金：钱记上、货没到，进度与状态都不动。"""
    _add(db, item, 0, 1000, deposit=True)
    item = _reload(db, item.id)

    assert compute.item_paid(item) == 1000      # 已付要有
    assert compute.item_paid_qty(item) == 0     # 已到货数量仍是 0
    assert compute.item_status(item) == "deposit"   # 状态是「已付定」，不是买完、也不是未买
    assert compute.item_dict(item)["bought"] is False


def test_deposit_reduces_both_unpaid_figures(db, item):
    """定金抵扣未付：两个口径各减已付，而不是一分不动。"""
    _add(db, item, 0, 1000, deposit=True)
    item = _reload(db, item.id)

    # 原价口径：6999 − 1000；日常价口径：5099 − 1000
    assert compute.item_unpaid(item) == 5999
    assert compute.item_daily_unpaid(item) == 4099


def test_balance_payment_finishes_the_item(db, item):
    """定金 1000 + 尾款 4099：结清、未付归零，优惠照旧算得出来。"""
    _add(db, item, 0, 1000, deposit=True)
    _add(db, item, 1, 4099, deposit=False)
    item = _reload(db, item.id)

    assert compute.item_paid(item) == 5099
    assert compute.item_paid_qty(item) == 1
    assert compute.item_status(item) == "done"
    assert compute.item_unpaid(item) == 0
    assert compute.item_daily_unpaid(item) == 0
    # 相对原价省了 6999 − 5099 = 1900（未付归零后这笔优惠要能显示出来）
    assert compute.item_actual_discount(item) == 1900


def test_deposit_larger_than_remaining_never_goes_negative(db, item):
    """预付超过未到货部分的价：未付压到 0，不出现负数。"""
    _add(db, item, 0, 8000, deposit=True)
    item = _reload(db, item.id)

    assert compute.item_unpaid(item) == 0
    assert compute.item_daily_unpaid(item) == 0


def test_partial_delivery_with_deposit(db, item):
    """到货一半、定金已付：未到货部分按价算，再减掉定金。"""
    _add(db, item, 0.5, 1000, deposit=True)
    _add(db, item, 0.5, 2550, deposit=False)
    item = _reload(db, item.id)

    assert compute.item_status(item) == "partial"
    # 未到货 0.5 台 × 6999 = 3499.5，减去定金 1000
    assert compute.item_unpaid(item) == 2499.5
    # 日常价口径：0.5 × 5099 = 2549.5 − 1000
    assert compute.item_daily_unpaid(item) == 1549.5


def test_excel_roundtrip_keeps_deposit_flag(db, item):
    """导出带「定金」列，回灌后还是定金（不是普通采购）。"""
    _add(db, item, 0, 1000, deposit=True)

    data = excel_io.export_xlsx(db, item.list_id)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    header = [c.value for c in wb["采购记录"][1]]
    assert wb["采购记录"].cell(row=2, column=header.index("定金") + 1).value == "是"

    excel_io.import_template(db, data, mode="replace", list_id=item.list_id)
    db.expire_all()
    record = db.query(PurchaseRecord).one()
    assert record.is_deposit is True

    reloaded = db.query(Item).filter(Item.list_id == item.list_id).one()
    assert compute.item_unpaid(reloaded) == 5999   # 抵扣效果跟着回来了


def test_old_file_without_deposit_column_still_imports(db, item):
    """老文件没有「定金」列：按普通采购读，行为与从前一致。"""
    _add(db, item, 1, 5099, deposit=False)
    data = excel_io.export_xlsx(db, item.list_id)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["采购记录"]
    header = [c.value for c in ws[1]]
    ws.delete_cols(header.index("定金") + 1)
    buf = io.BytesIO()
    wb.save(buf)

    report = excel_io.import_template(db, buf.getvalue(), mode="replace",
                                      list_id=item.list_id)
    assert report.get("warnings") == [], report.get("warnings")
    db.expire_all()
    assert db.query(PurchaseRecord).one().is_deposit is False


def test_sync_payload_carries_deposit_flag(db, item):
    """同步载荷带上定金：不然桌面端与手机端一合并就把它抹平了。"""
    from app.services import list_transfer
    _add(db, item, 0, 1000, deposit=True)

    lst = db.query(ItemList).filter(ItemList.id == item.list_id).one()
    payload = list_transfer.export_list(db, lst)
    record = payload["items"][0]["records"][0]
    assert record["is_deposit"] is True

    # 合并结果里也要留着（基线就是服务器那份）
    from app.services import merge
    base = merge.Baseline(payload=payload,
                          local_map=payload.get("local_map") or {})
    result = merge.merge(base, payload, payload)
    assert result.payload["items"][0]["records"][0]["is_deposit"] is True


def test_item_put_keeps_deposit_flag(client):
    """界面保存物料时走的是 PUT /api/items（整条替换记录）—— 定金字要能存住。

    这条是补的回归：`_apply_records` 里漏写 is_deposit 时，勾了定金等于没勾
    （钱算了、进度也照涨），界面上表现为「复选框没生效」。
    """
    created = client.post("/api/items", json={
        "name": "定金回归", "unit": "台", "qty_total": 1, "price": 6999,
        "discount_price": 5099, "allocations": [], "records": [],
    }, headers=_hdr(client)).json()

    saved = client.put(f"/api/items/{created['id']}", json={
        "name": "定金回归", "unit": "台", "qty_total": 1, "price": 6999,
        "discount_price": 5099, "allocations": [], "base_rev": created["rev"],
        "records": [{"qty": 0, "amount": 1000, "is_deposit": True,
                     "date": "2026-10-08", "room_ids": []}],
    }, headers=_hdr(client)).json()

    assert saved["records"][0]["is_deposit"] is True
    assert saved["unpaid"] == 5999          # 定金抵扣了未付
    assert saved["daily_unpaid"] == 4099
    assert saved["status"] == "deposit"     # 货没到，状态是「已付定」
    assert saved["paid_qty"] == 0


def test_summary_counts_deposit_and_lists_it_as_pending(client):
    """总览：已付定单独计数，而且照样进「未采购清单」（钱付了、东西还得盯着）。"""
    created = client.post("/api/items", json={
        "name": "付了定金的", "unit": "台", "qty_total": 1, "price": 6999,
        "discount_price": 5099, "allocations": [], "records": [],
    }, headers=_hdr(client)).json()
    client.put(f"/api/items/{created['id']}", json={
        "name": "付了定金的", "unit": "台", "qty_total": 1, "price": 6999,
        "discount_price": 5099, "allocations": [], "base_rev": created["rev"],
        "records": [{"qty": 0, "amount": 1000, "is_deposit": True, "room_ids": []}],
    }, headers=_hdr(client))

    data = client.get("/api/summary", headers=_hdr(client)).json()
    assert data["totals"]["status_count"]["deposit"] == 1
    assert data["totals"]["status_count"]["unbought"] == 0
    assert [v["name"] for v in data["unbought"]] == ["付了定金的"]


def test_sync_upload_keeps_deposit_flag(db, item):
    """手机整份上传（import_list 落库）也要留住定金 —— 这是会丢数据的写路径。"""
    from app.services import list_transfer
    _add(db, item, 0, 1000, deposit=True)

    lst = db.query(ItemList).filter(ItemList.id == item.list_id).one()
    payload = list_transfer.export_list(db, lst)

    # 换一份新清单整份灌进去（模拟手机首次上传）
    list_transfer.import_list(db, payload)

    db.expire_all()
    other = db.query(ItemList).filter(ItemList.name == lst.name,
                                      ItemList.id != lst.id).one()
    record = (db.query(PurchaseRecord)
              .join(Item).filter(Item.list_id == other.id).one())
    assert record.is_deposit is True
    reloaded = record.item
    assert compute.item_unpaid(reloaded) == 5999   # 抵扣效果跟着过来了


def test_undo_restores_deposit_record(db, item):
    """回退一次「记定金」：那条定金记录要原样回来，标记也在。"""
    from app import audit
    before = audit.snapshot_item(db, item.id)
    _add(db, item, 0, 1000, deposit=True)
    assert db.query(PurchaseRecord).count() == 1

    audit._restore_item(db, before)
    db.expire_all()
    assert db.query(PurchaseRecord).count() == 0

    # 反过来：有定金 → 快照留下 → 恢复后仍是定金
    _add(db, item, 0, 1000, deposit=True)
    snap = audit.snapshot_item(db, item.id)
    db.query(PurchaseRecord).delete()
    db.commit()
    audit._restore_item(db, snap)
    db.expire_all()
    assert db.query(PurchaseRecord).one().is_deposit is True
