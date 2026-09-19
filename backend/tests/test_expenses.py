"""额外费用（运费 / 安装费 / 辅料）。

它存在的意义是把"运费摊进单价"这个老做法替掉：单价回归真实，
"实际优惠"也不会因为摊了钱而算成负数。代价是它独立于两个金额口径 ——
所以最要紧的一条测试是：**加多少费用，物料的三个金额数字都不许动**。
"""

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import auth
from app.db import SessionLocal, engine
from app.main import app
from app.models import (Allocation, Category, ExtraExpense, Item, ItemList,
                        PurchaseRecord, Room, User)
from app.seed import init_db
from app.services import excel_io
from tests.conftest import TEST_PASSWORD as PASSWORD, TEST_USER as USER


@pytest.fixture()
def client():
    init_db()
    session = SessionLocal()
    for table in (ExtraExpense, Allocation, Item, PurchaseRecord, Room, Category,
                  User, ItemList):
        session.query(table).delete()
    session.commit()
    session.add(ItemList(name="采购清单", sort=0))
    session.commit()
    session.close()
    auth._failures.clear()
    with TestClient(app) as c:
        c.post("/api/auth/setup", json={"username": USER, "password": PASSWORD})
        yield c
    engine.dispose()


def _hdr(client):
    return {"X-List-Id": str(client.get("/api/lists").json()[0]["id"])}


def _mk_item(client, name="筒灯", qty=4, price=100):
    return client.post("/api/items", json={"name": name, "unit": "个",
                                           "qty_total": qty, "price": price},
                       headers=_hdr(client)).json()


def test_expenses_do_not_touch_the_money_breakdown(client):
    """最要紧的一条：加多少运费，两个口径的所有数字一字不变。"""
    item = _mk_item(client)
    client.post(f"/api/items/{item['id']}/records", json={"qty": 1, "amount": 80},
                headers=_hdr(client))
    before = client.get("/api/summary", headers=_hdr(client)).json()["totals"]

    r = client.post("/api/expenses", json={"kind": "运费", "amount": 60,
                                           "date": "2026-09-16", "vendor": "京东",
                                           "note": "整单运费"},
                    headers=_hdr(client))
    assert r.status_code == 200, r.text

    data = client.get("/api/summary", headers=_hdr(client)).json()
    assert data["totals"] == before           # 口径一个字段都没变
    assert data["expenses_total"] == 60.0
    assert data["expenses_by_kind"] == [{"kind": "运费", "amount": 60.0}]
    assert data["expenses_count"] == 1
    # 物料自己的数字也没被带动
    got = client.get(f"/api/items/{item['id']}", headers=_hdr(client)).json()
    assert got["paid"] == 80 and got["status"] == "partial"


def test_expenses_crud_and_kinds(client):
    a = client.post("/api/expenses", json={"kind": "运费", "amount": 60},
                    headers=_hdr(client)).json()
    b = client.post("/api/expenses", json={"kind": "安装费", "amount": 300},
                    headers=_hdr(client)).json()
    client.post("/api/expenses", json={"kind": "搬运费", "amount": 100},
                headers=_hdr(client))  # 名目不受限，自己写也行

    data = client.get("/api/summary", headers=_hdr(client)).json()
    assert data["expenses_total"] == 460.0
    assert {k["kind"]: k["amount"] for k in data["expenses_by_kind"]} == {
        "运费": 60.0, "安装费": 300.0, "搬运费": 100.0}

    upd = client.put(f"/api/expenses/{a['id']}",
                     json={"kind": "运费", "amount": 88, "note": "改过"},
                     headers=_hdr(client))
    assert upd.status_code == 200 and upd.json()["amount"] == 88

    assert client.delete(f"/api/expenses/{b['id']}",
                         headers=_hdr(client)).status_code == 200
    assert client.get("/api/summary", headers=_hdr(client)).json()["expenses_total"] == 188.0


def test_expense_date_is_validated_and_normalized(client):
    ok = client.post("/api/expenses", json={"kind": "运费", "amount": 10,
                                            "date": "2026/9/4"}, headers=_hdr(client))
    assert ok.status_code == 200 and ok.json()["date"] == "2026-09-04"
    bad = client.post("/api/expenses", json={"kind": "运费", "amount": 10,
                                             "date": "上周三"}, headers=_hdr(client))
    assert bad.status_code == 422


def test_expenses_are_isolated_between_lists(client):
    first = client.get("/api/lists").json()[0]
    other = client.post("/api/lists", json={"name": "年货清单"}).json()
    hdr_other = {"X-List-Id": str(other["id"])}

    client.post("/api/expenses", json={"kind": "运费", "amount": 60},
                headers=_hdr(client))
    assert client.get("/api/expenses", headers=_hdr(client)).json()[0]["amount"] == 60
    assert client.get("/api/expenses", headers=hdr_other).json() == []
    assert client.get("/api/summary", headers=hdr_other).json()["expenses_total"] == 0

    # 别的清单里改不到它
    eid = client.get("/api/expenses", headers=_hdr(client)).json()[0]["id"]
    assert client.delete(f"/api/expenses/{eid}", headers=hdr_other).status_code == 404
    assert client.get("/api/expenses", headers=_hdr(client)).json() != []


def test_expense_can_link_to_an_item(client):
    item = _mk_item(client, "筒灯")
    r = client.post("/api/expenses", json={"kind": "运费", "amount": 60,
                                           "item_id": item["id"]}, headers=_hdr(client))
    assert r.status_code == 200, r.text
    assert r.json()["item_id"] == item["id"]
    assert r.json()["item_name"] == "筒灯"

    row = client.get("/api/expenses", headers=_hdr(client)).json()[0]
    assert row["item_name"] == "筒灯"
    # 关联不影响口径
    assert client.get("/api/summary", headers=_hdr(client)).json()["expenses_total"] == 60


def test_expense_item_must_belong_to_current_list(client):
    other = client.post("/api/lists", json={"name": "年货清单"}).json()
    foreign = client.post("/api/items",
                          json={"name": "瓜子", "unit": "包", "qty_total": 1, "price": 1},
                          headers={"X-List-Id": str(other["id"])}).json()
    r = client.post("/api/expenses", json={"kind": "运费", "amount": 10,
                                           "item_id": foreign["id"]}, headers=_hdr(client))
    assert r.status_code == 400 and "不属于当前清单" in r.json()["detail"]


def test_expense_survives_item_deletion(client):
    """删物料不能让费用记录消失 —— 那是钱，只是关联断掉。"""
    item = _mk_item(client, "筒灯")
    client.post("/api/expenses", json={"kind": "运费", "amount": 60,
                                       "item_id": item["id"]}, headers=_hdr(client))

    client.delete(f"/api/items/{item['id']}", headers=_hdr(client))  # 移入回收站
    row = client.get("/api/expenses", headers=_hdr(client)).json()[0]
    assert row["amount"] == 60 and row["item_name"] == "筒灯"

    # 彻底删掉物料：关联断掉，费用还在
    client.delete(f"/api/trash/{item['id']}", headers=_hdr(client))
    row = client.get("/api/expenses", headers=_hdr(client)).json()[0]
    assert row["amount"] == 60 and row["item_id"] is None and row["item_name"] == ""


def test_excel_roundtrip_keeps_expenses(client):
    session = SessionLocal()
    try:
        item = Item(name="筒灯", unit="个", qty_total=4, price=100, list_id=1)
        session.add(item)
        session.commit()
    finally:
        session.close()
    client.post("/api/expenses", json={"kind": "运费", "amount": 60,
                                       "date": "2026-09-14", "vendor": "京东",
                                       "order_no": "JD001"}, headers=_hdr(client))

    exported = client.get("/api/export", headers=_hdr(client)).content
    ws = openpyxl.load_workbook(io.BytesIO(exported))["额外费用"]
    assert [c.value for c in ws[1]] == ["类型", "金额", "日期", "商家", "订单号", "备注"]

    # 先清空再回灌：费用得回到 60
    client.delete(f"/api/expenses/{client.get('/api/expenses', headers=_hdr(client)).json()[0]['id']}",
                  headers=_hdr(client))
    assert client.get("/api/summary", headers=_hdr(client)).json()["expenses_total"] == 0
    r = client.post("/api/import",
                    files={"file": ("a.xlsx", io.BytesIO(exported),
                                    "application/vnd.openxmlformats-officedocument"
                                    ".spreadsheetml.sheet")},
                    data={"mode": "replace"}, headers=_hdr(client))
    assert r.status_code == 200, r.text
    assert r.json()["expenses"] == 1
    items = client.get("/api/expenses", headers=_hdr(client)).json()
    assert items[0]["amount"] == 60 and items[0]["vendor"] == "京东"


def test_old_file_without_expense_sheet_keeps_existing(client):
    """老备份里没有「额外费用」这一页：导入不能把现有的费用清掉。"""
    client.post("/api/expenses", json={"kind": "运费", "amount": 60},
                headers=_hdr(client))
    exported = client.get("/api/export", headers=_hdr(client)).content
    wb = openpyxl.load_workbook(io.BytesIO(exported))
    wb.remove(wb["额外费用"])          # 造一份"老版本导出"的文件
    buf = io.BytesIO()
    wb.save(buf)

    r = client.post("/api/import",
                    files={"file": ("old.xlsx", buf.getvalue(),
                                    "application/vnd.openxmlformats-officedocument"
                                    ".spreadsheetml.sheet")},
                    data={"mode": "replace"}, headers=_hdr(client))
    assert r.status_code == 200, r.text
    assert r.json()["expenses"] == 0
    assert client.get("/api/summary", headers=_hdr(client)).json()["expenses_total"] == 60.0


def test_deleting_list_removes_its_expenses(client):
    first = client.get("/api/lists").json()[0]
    other = client.post("/api/lists", json={"name": "年货清单"}).json()
    client.post("/api/expenses", json={"kind": "运费", "amount": 60}, headers=_hdr(client))
    client.post("/api/expenses", json={"kind": "运费", "amount": 20},
                headers={"X-List-Id": str(other["id"])})

    assert client.delete(f"/api/lists/{first['id']}").status_code == 200
    session = SessionLocal()
    try:
        left = session.query(ExtraExpense).all()
        assert len(left) == 1 and left[0].amount == 20
    finally:
        session.close()
