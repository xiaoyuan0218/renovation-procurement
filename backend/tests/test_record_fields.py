"""采购记录上的「商家 / 订单号」：选填字段，只做记录，不参与任何金额口径。

重点是两件事：写进去能读回来（接口 + Excel 往返），以及老文件／老备份里
没有这两列时不能出问题。
"""

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import auth
from app.db import SessionLocal, engine
from app.main import app
from app.models import (Allocation, Category, Item, ItemList, PurchaseRecord,
                        Room, User)
from app.seed import init_db
from app.services import compute, excel_io
from tests.conftest import TEST_PASSWORD as PASSWORD, TEST_USER as USER


@pytest.fixture()
def db():
    init_db()
    session = SessionLocal()
    for table in (Allocation, Item, PurchaseRecord, Room, Category, User, ItemList):
        session.query(table).delete()
    session.commit()
    session.add(ItemList(name="采购清单", sort=0))
    session.commit()
    yield session
    session.close()
    engine.dispose()


@pytest.fixture()
def client():
    init_db()
    session = SessionLocal()
    for table in (Allocation, Item, PurchaseRecord, Room, Category, User, ItemList):
        session.query(table).delete()
    session.commit()
    session.add(ItemList(name="采购清单", sort=0))
    session.commit()
    session.close()
    auth._failures.clear()
    with TestClient(app) as c:
        c.post("/api/auth/setup", json={"username": USER, "password": PASSWORD})
        yield c
    engine.dispose()


def test_api_roundtrip_vendor_and_order_no(client):
    item = client.post("/api/items", json={"name": "筒灯", "unit": "个",
                                           "qty_total": 4, "price": 99}).json()
    r = client.post(f"/api/items/{item['id']}/records", json={
        "qty": 2, "amount": 180, "date": "2026-09-14",
        "vendor": "京东", "order_no": "JD20260914001", "note": "第一批",
    })
    assert r.status_code == 200, r.text
    rec = r.json()["records"][0]
    assert rec["vendor"] == "京东"
    assert rec["order_no"] == "JD20260914001"

    # 单条记录也能单独改这两个字段
    upd = client.put(f"/api/records/{rec['id']}", json={"vendor": "天猫"})
    assert upd.status_code == 200
    assert upd.json()["records"][0]["vendor"] == "天猫"
    assert upd.json()["records"][0]["order_no"] == "JD20260914001"

    # 整条物料保存时也带得回来
    detail = client.get(f"/api/items/{item['id']}").json()
    saved = client.put(f"/api/items/{item['id']}", json={
        "name": detail["name"], "unit": detail["unit"],
        "qty_total": detail["qty_total"], "price": detail["price"],
        "allocations": [], "base_rev": detail["rev"],
        "records": [{"qty": 2, "amount": 180, "date": "2026-09-14",
                     "vendor": "天猫", "order_no": "JD20260914001", "note": "第一批"}],
    })
    assert saved.status_code == 200, saved.text
    assert saved.json()["records"][0]["vendor"] == "天猫"


def test_patch_with_null_fields_leaves_them_alone(client):
    """只改一个字段时，其它字段哪怕显式传了 null 也不许被清空。

    Kotlin 客户端会把整个对象序列化出去（日期、备注都是 null），
    这条就是那类客户端的兜底。
    """
    item = client.post("/api/items", json={"name": "筒灯", "unit": "个",
                                           "qty_total": 4, "price": 99}).json()
    rec = client.post(f"/api/items/{item['id']}/records", json={
        "qty": 2, "amount": 180, "date": "2026-09-14",
        "vendor": "京东", "order_no": "JD001", "note": "第一批",
    }).json()["records"][0]

    r = client.put(f"/api/records/{rec['id']}", json={
        "amount": 200, "qty": None, "date": None, "note": None,
        "vendor": None, "order_no": None,
    })
    assert r.status_code == 200, r.text
    got = r.json()["records"][0]
    assert got["amount"] == 200
    assert got["qty"] == 2 and got["date"] == "2026-09-14"
    assert got["vendor"] == "京东" and got["order_no"] == "JD001"
    assert got["note"] == "第一批"


def test_amounts_ignore_vendor_fields(client):
    """加了两个字段不能碰到金额：同样的数量金额，填不填商家结果一样。"""
    a = client.post("/api/items", json={"name": "带商家", "unit": "个",
                                        "qty_total": 2, "price": 100}).json()
    b = client.post("/api/items", json={"name": "不带商家", "unit": "个",
                                        "qty_total": 2, "price": 100}).json()
    client.post(f"/api/items/{a['id']}/records",
                json={"qty": 1, "amount": 80, "vendor": "京东", "order_no": "X1"})
    client.post(f"/api/items/{b['id']}/records", json={"qty": 1, "amount": 80})
    got_a = client.get(f"/api/items/{a['id']}").json()
    got_b = client.get(f"/api/items/{b['id']}").json()
    for key in ("list_total", "discount_total", "paid", "unpaid", "status"):
        assert got_a[key] == got_b[key], key


def test_excel_roundtrip_keeps_vendor_fields(db):
    room = Room(name="客厅", sort=0, list_id=1)
    db.add(room)
    db.flush()
    item = Item(name="筒灯", price=99, unit="个", qty_total=4, list_id=1)
    db.add(item)
    db.flush()
    db.add(Allocation(item_id=item.id, room_id=room.id, qty=4))
    item.records.append(PurchaseRecord(qty=2, amount=180, date="2026-09-14",
                                       vendor="京东", order_no="JD001", note="第一批"))
    db.commit()

    data = excel_io.export_xlsx(db, list_id=1)
    ws = openpyxl.load_workbook(io.BytesIO(data))["采购记录"]
    header = [c.value for c in ws[1]]
    row = next(ws.iter_rows(min_row=2, values_only=True))
    rec = dict(zip(header, row))
    assert rec["商家"] == "京东" and rec["订单号"] == "JD001"

    # 覆盖导入后仍然是这两个值
    excel_io.import_template(db, data, mode="replace", list_id=1)
    db.expire_all()
    saved = db.query(PurchaseRecord).one()
    assert (saved.vendor, saved.order_no) == ("京东", "JD001")
    # 金额口径不受影响
    assert compute.item_paid(db.query(Item).one()) == 180


def test_import_file_without_vendor_columns_still_works(db):
    """老导出文件没有这两列：导入照常，两字段留空。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "物料汇总"
    ws.append(["类目", "物料名称", "单位", "数量", "单价", "优惠单价",
               "日常价", "实付数量", "实付金额", "未付数量", "未付金额",
               "已购", "备注"])
    ws.append(["照明", "筒灯", "个", 2, 99, None, None, None, None, None, None, "否", ""])
    ws3 = wb.create_sheet("采购记录")
    ws3.append(["物料名称", "实付数量", "实付金额", "付款日期", "备注", "物料ID"])
    ws3.append(["筒灯", 2, 180, "2026-09-14", "老文件", None])
    buf = io.BytesIO()
    wb.save(buf)

    report = excel_io.import_template(db, buf.getvalue(), mode="replace")
    assert report["records"] == 1
    db.expire_all()
    saved = db.query(PurchaseRecord).one()
    assert (saved.vendor, saved.order_no) == ("", "")
    assert saved.note == "老文件"
