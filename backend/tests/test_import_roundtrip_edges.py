"""导入导出的边界往返：这些是 PC 与手机之间传数据的唯一通道，错一次就是数据错。

主往返测试在 test_import_template.py（同名物料、幂等、覆盖价、老文件）。这里补
的是那边没覆盖、但真实数据里会踩到的几类：分组名含分隔符、空分组、多分组记录、
特殊字符、浮点精度、坏数字提示，以及「模板与导出逐页一致」。
"""

import io

import openpyxl
import pytest

from app.db import SessionLocal, engine
from app.models import (Allocation, Category, Item, ItemList, PurchaseRecord,
                        RecordRoom, Room)
from app.seed import init_db
from app.services import codes, compute, excel_io


@pytest.fixture()
def db():
    init_db()
    session = SessionLocal()
    for table in (RecordRoom, Allocation, PurchaseRecord, Item, Room, Category):
        session.query(table).delete()
    session.commit()
    yield session
    session.close()
    engine.dispose()


def _first_list(db) -> ItemList:
    return db.query(ItemList).order_by(ItemList.sort, ItemList.id).first()


def _new_list(db, name: str) -> ItemList:
    lst = ItemList(name=name, sort=99, code=codes.new_code())
    db.add(lst)
    db.commit()
    return lst


def _rooms_of(db, list_id) -> set:
    db.expire_all()
    return {r.name for r in db.query(Room).filter(Room.list_id == list_id).all()}


def _cats_of(db, list_id) -> set:
    db.expire_all()
    return {c.name for c in db.query(Category).filter(Category.list_id == list_id).all()}


def _snapshot(db) -> dict:
    """所有会进金额口径的量 + 分组/分类/明细，用来判断一次导入有没有改动任何东西。

    刻意**只按名字对账，不碰任何行 id**：落地是「清掉重插」，id 本来就会重排，
    拿 id 比会把「数据没变」误判成「变了」。
    """
    db.expire_all()
    room_name = {r.id: r.name for r in db.query(Room).all()}
    item_name = {i.id: i.name for i in db.query(Item).all()}
    views = [compute.item_dict(i) for i in db.query(Item).all()]
    records = {r.id: r for r in db.query(PurchaseRecord).all()}
    return {
        "items": db.query(Item).count(),
        "rooms": sorted(room_name.values()),
        "cats": sorted(c.name for c in db.query(Category).all()),
        "list_total": round(sum(v["list_total"] for v in views), 6),
        "discount_total": round(sum(v["discount_total"] for v in views), 6),
        "paid": round(sum(v["paid"] for v in views), 6),
        "unpaid": round(sum(v["unpaid"] for v in views), 6),
        "per_item": sorted((v["name"], v["total_qty"], v["list_total"],
                            v["paid"], v["unpaid"]) for v in views),
        "allocs": sorted((item_name.get(a.item_id, ""), room_name.get(a.room_id, ""),
                          a.qty, a.price_override, a.note or "")
                         for a in db.query(Allocation).all()),
        "records": sorted((item_name.get(r.item_id, ""), r.qty, r.amount, r.date or "",
                           r.note or "", r.vendor or "", r.order_no or "")
                          for r in db.query(PurchaseRecord).all()),
        "record_rooms": sorted(
            (item_name.get(records[rr.record_id].item_id, ""), room_name.get(rr.room_id, ""))
            for rr in db.query(RecordRoom).all() if rr.record_id in records),
    }


# ---------------------------------------------------------------- 分组名含分隔符

def test_room_name_with_slash_survives_cross_list(db):
    """分组名叫「客厅/餐厅」，搬到另一份清单后不能被拆成「客厅」和「餐厅」。

    采购记录的「分组」列用顿号连接多个分组，导入时按 `、,，/` 拆 —— 名字本身
    含这些字符的就会被拆坏。修法是导出时把分组名写进「分组」页，导入时整串能
    匹配上就不拆。
    """
    src = _first_list(db)
    room = Room(name="客厅/餐厅", sort=0, list_id=src.id)
    db.add(room)
    db.flush()
    item = Item(name="筒灯", unit="个", qty_total=2, price=10, list_id=src.id)
    db.add(item)
    db.flush()
    db.add(Allocation(item_id=item.id, room_id=room.id, qty=2))
    record = PurchaseRecord(qty=2, amount=20)
    record.rooms.append(RecordRoom(room_id=room.id))
    item.records.append(record)
    db.commit()

    data = excel_io.export_xlsx(db, src.id)
    dst = _new_list(db, "目标清单")
    excel_io.import_template(db, data, mode="replace", list_id=dst.id)

    assert _rooms_of(db, dst.id) == {"客厅/餐厅"}


def test_room_name_with_dunhao_survives_cross_list(db):
    """顿号同理：「客厅、过道」是一个分组名，不是一个分组加另一个。"""
    src = _first_list(db)
    room = Room(name="客厅、过道", sort=0, list_id=src.id)
    db.add(room)
    db.flush()
    item = Item(name="筒灯", unit="个", qty_total=1, price=10, list_id=src.id)
    db.add(item)
    db.flush()
    db.add(Allocation(item_id=item.id, room_id=room.id, qty=1))
    db.commit()

    data = excel_io.export_xlsx(db, src.id)
    dst = _new_list(db, "目标清单")
    excel_io.import_template(db, data, mode="replace", list_id=dst.id)

    assert _rooms_of(db, dst.id) == {"客厅、过道"}


# ---------------------------------------------------------------- 空分组 / 空分类

def test_empty_rooms_and_categories_travel_along(db):
    """没有物料的空分组、空分类也要跟着走 —— 它们只存在于「分组」「分类」页里。"""
    src = _first_list(db)
    db.add_all([
        Room(name="客厅", sort=0, list_id=src.id),
        Room(name="还没买过东西", sort=1, list_id=src.id),
        Category(name="照明", sort=0, list_id=src.id),
        Category(name="备用分类", sort=1, list_id=src.id),
    ])
    db.commit()

    data = excel_io.export_xlsx(db, src.id)
    dst = _new_list(db, "目标清单")
    excel_io.import_template(db, data, mode="replace", list_id=dst.id)

    assert _rooms_of(db, dst.id) == {"客厅", "还没买过东西"}
    assert _cats_of(db, dst.id) == {"照明", "备用分类"}


# ---------------------------------------------------------------- 多分组采购记录

def test_multi_room_record_roundtrip(db):
    """一笔付款同时买了两间的东西，Excel 往返后归属两间都要留着。"""
    src = _first_list(db)
    r1 = Room(name="客厅", sort=0, list_id=src.id)
    r2 = Room(name="主卧", sort=1, list_id=src.id)
    db.add_all([r1, r2])
    db.flush()
    item = Item(name="筒灯", unit="个", qty_total=4, price=30, list_id=src.id)
    db.add(item)
    db.flush()
    record = PurchaseRecord(qty=4, amount=120)
    record.rooms.append(RecordRoom(room_id=r1.id))
    record.rooms.append(RecordRoom(room_id=r2.id))
    item.records.append(record)
    db.commit()

    before = _snapshot(db)
    data = excel_io.export_xlsx(db, src.id)
    excel_io.import_template(db, data, mode="merge", list_id=src.id)
    assert _snapshot(db) == before

    # 归属要能按名字还原（行 id 会被重排，不能直接比 id）
    db.expire_all()
    rec = db.query(PurchaseRecord).one()
    assert {db.get(Room, rr.room_id).name for rr in rec.rooms} == {"客厅", "主卧"}


# ---------------------------------------------------------------- 空清单

def test_empty_list_roundtrip(db):
    """空清单导出再导入：什么都不该被造出来，也不该报错。"""
    src = _first_list(db)
    data = excel_io.export_xlsx(db, src.id)
    report = excel_io.import_template(db, data, mode="replace", list_id=src.id)
    assert report["items_created"] == 0
    assert report["items_matched"] == 0
    assert db.query(Item).count() == 0


# ---------------------------------------------------------------- 特殊字符

def test_special_characters_roundtrip(db):
    """换行、引号、尖括号、& 这些在 XML 里有特殊含义的字符必须原样往返。"""
    src = _first_list(db)
    weird = '带"双引号"和\'单引号\'\n还有换行 & <尖括号>'
    item = Item(name="特殊物料", note=weird, brand=weird, model="A<B>&C",
                unit="个", qty_total=1, price=1, list_id=src.id)
    db.add(item)
    db.commit()

    data = excel_io.export_xlsx(db, src.id)
    excel_io.import_template(db, data, mode="replace", list_id=src.id)

    db.expire_all()
    got = db.query(Item).filter(Item.name == "特殊物料").one()
    assert got.note == weird
    assert got.brand == weird
    assert got.model == "A<B>&C"


def test_long_text_roundtrip(db):
    """备注接近字段上限时也要完整往返（导出不能截断）。"""
    src = _first_list(db)
    long_note = "备" * 480
    db.add(Item(name="长备注物料", note=long_note, unit="个", qty_total=1,
                price=1, list_id=src.id))
    db.commit()

    data = excel_io.export_xlsx(db, src.id)
    excel_io.import_template(db, data, mode="replace", list_id=src.id)

    db.expire_all()
    assert db.query(Item).filter(Item.name == "长备注物料").one().note == long_note


# ---------------------------------------------------------------- 浮点精度

def test_float_precision_roundtrip(db):
    """金额往返不丢精度。

    xlsx（Excel）本身只保留 15 位有效数字，超出的部分会被舍入 —— 这是格式的
    固有限制，不是这一层能修的（实测 0.1+0.2 那种 17 位尾数写进去会变成 0.3）。
    真实金额远在 15 位以内，这条盯的是常见写法（0.1、0.07、三位小数、带小数的
    单价×数量）能分毫不差地回来。
    """
    src = _first_list(db)
    item = Item(name="精密件", unit="个", qty_total=3, price=0.1,
                discount_price=0.07, list_id=src.id)
    db.add(item)
    db.flush()
    item.records.append(PurchaseRecord(qty=3, amount=0.3))
    item.records.append(PurchaseRecord(qty=1, amount=1234.567))
    db.commit()

    before = _snapshot(db)
    data = excel_io.export_xlsx(db, src.id)
    excel_io.import_template(db, data, mode="replace", list_id=src.id)
    assert _snapshot(db) == before


# ---------------------------------------------------------------- 坏数字要给提示

def test_bad_number_produces_warning(db):
    """数量填成文本不该被静默当 0 —— 用户得知道哪一行没读进去。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "物料汇总"
    ws.append(["类目", "物料名称", "单位", "数量", "单价"])
    ws.append(["照明", "筒灯", "个", "十个", 30])
    buf = io.BytesIO()
    wb.save(buf)

    report = excel_io.import_template(db, buf.getvalue(), mode="replace")
    assert any("数量" in w and "不是数字" in w for w in report["warnings"]), report["warnings"]


def test_bad_number_in_record_produces_warning(db):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "物料汇总"
    ws.append(["类目", "物料名称", "单位", "数量", "单价"])
    ws.append(["照明", "筒灯", "个", 2, 30])
    ws2 = wb.create_sheet("采购记录")
    ws2.append(["物料名称", "实付数量", "实付金额", "付款日期"])
    ws2.append(["筒灯", "两", 60, "2026-09-01"])
    buf = io.BytesIO()
    wb.save(buf)

    report = excel_io.import_template(db, buf.getvalue(), mode="replace")
    assert any("实付数量" in w and "不是数字" in w for w in report["warnings"]), report["warnings"]


def test_blank_cells_do_not_warn(db):
    """没填就是没填，不该被当成「填错了」刷屏。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "物料汇总"
    ws.append(["类目", "物料名称", "单位", "数量", "单价", "优惠单价", "实付数量", "实付金额"])
    ws.append(["照明", "筒灯", "个", 2, 30, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)

    report = excel_io.import_template(db, buf.getvalue(), mode="replace")
    assert report["warnings"] == [], report["warnings"]


# ---------------------------------------------------------------- 模板与导出一致

def test_template_header_matches_export_on_every_sheet(db):
    """模板的每一页表头都要和导出逐列一致 —— 用户会把导出文件当模板用。"""
    tpl = openpyxl.load_workbook(io.BytesIO(excel_io.build_template()))
    exp = openpyxl.load_workbook(io.BytesIO(excel_io.export_xlsx(db)))
    for sheet in ("物料汇总", "布点明细", "采购记录", "额外费用", "分组", "分类"):
        assert sheet in tpl.sheetnames, f"模板缺少「{sheet}」页"
        assert sheet in exp.sheetnames, f"导出缺少「{sheet}」页"
        assert [c.value for c in tpl[sheet][1]] == [c.value for c in exp[sheet][1]], sheet


def test_export_reimport_is_noop_with_new_sheets(db):
    """加了「分组」「分类」页之后，导出→回灌仍然是零变化。"""
    src = _first_list(db)
    r1 = Room(name="客厅", sort=0, list_id=src.id)
    r2 = Room(name="还没用过", sort=1, list_id=src.id)
    db.add_all([r1, r2])
    cat = Category(name="照明", sort=0, list_id=src.id)
    db.add(cat)
    db.flush()
    item = Item(name="筒灯", unit="个", qty_total=6, price=30, discount_price=28,
                category_id=cat.id, list_id=src.id)
    db.add(item)
    db.flush()
    db.add(Allocation(item_id=item.id, room_id=r1.id, qty=6))
    item.records.append(PurchaseRecord(qty=6, amount=180, date="2026-09-01"))
    db.commit()

    before = _snapshot(db)
    for _ in range(2):  # 连灌两次也不能累积
        data = excel_io.export_xlsx(db, src.id)
        excel_io.import_template(db, data, mode="merge", list_id=src.id)
    assert _snapshot(db) == before
