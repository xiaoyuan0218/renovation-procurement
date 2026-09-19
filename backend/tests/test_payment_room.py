"""付款归属分组：让"哪间买齐了"从猜变成算。

老做法是拿累计实付按分配顺序往下抵 —— 先买了餐厅的灯，也会把排在前面的
客厅标成"已买齐"。现在付款时可以写明"这笔是给哪个分组花的"，写了就按它算。
"""

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import auth
from app.db import SessionLocal, engine
from app.main import app
from app.models import (Allocation, Category, ExtraExpense, Item, ItemList,
                        PurchaseRecord, Room, User)
from app.seed import init_db
from app.services import compute
from tests.conftest import TEST_PASSWORD as PASSWORD, TEST_USER as USER


@pytest.fixture()
def client():
    init_db()
    session = SessionLocal()
    for table in (ExtraExpense, Allocation, Item, PurchaseRecord, Room, Category,
                  User, ItemList):
        session.query(table).delete()
    session.commit()
    session.add(ItemList(name="采购清单", sort=0))
    session.commit()
    session.close()
    auth._failures.clear()
    with TestClient(app) as c:
        c.post("/api/auth/setup", json={"username": USER, "password": PASSWORD})
        yield c
    engine.dispose()


def _hdr(client):
    return {"X-List-Id": str(client.get("/api/lists").json()[0]["id"])}


def _scene(client, living_qty=3, dining_qty=4):
    """一个物料分给客厅和餐厅两个分组。"""
    living = client.post("/api/rooms", json={"name": "客厅", "sort": 0},
                         headers=_hdr(client)).json()
    dining = client.post("/api/rooms", json={"name": "餐厅", "sort": 1},
                         headers=_hdr(client)).json()
    item = client.post("/api/items", json={"name": "筒灯", "unit": "个", "price": 100},
                       headers=_hdr(client)).json()
    client.put("/api/items/" + str(item["id"]), headers=_hdr(client), json={
        "name": "筒灯", "unit": "个", "qty_total": 0, "price": 100,
        "allocations": [{"room_id": living["id"], "qty": living_qty},
                        {"room_id": dining["id"], "qty": dining_qty}],
        "records": [],
        "base_rev": client.get(f"/api/items/{item['id']}",
                               headers=_hdr(client)).json()["rev"],
    })
    return item, living, dining


def _cell(client, item_id, room_id):
    matrix = client.get("/api/matrix", headers=_hdr(client)).json()
    row = next(i for i in matrix["items"] if i["id"] == item_id)
    return row["cells"].get(str(room_id), {}).get("paid_qty", 0)


def test_paying_one_room_marks_only_that_room(client):
    """只买了餐厅（4 个，写明归属）：餐厅标绿，客厅一个都不标。

    这一条就是本功能的全部意义 —— 老算法会把客厅（排在前面）也标成已买齐。
    """
    item, living, dining = _scene(client)
    r = client.post(f"/api/items/{item['id']}/records",
                    json={"qty": 4, "amount": 400, "room_id": dining["id"]},
                    headers=_hdr(client))
    assert r.status_code == 200, r.text
    assert r.json()["records"][0]["room_id"] == dining["id"]

    assert _cell(client, item["id"], dining["id"]) == 4    # 餐厅：买齐了
    assert _cell(client, item["id"], living["id"]) == 0    # 客厅：没买，别瞎标


def test_multiple_payments_each_to_their_room(client):
    item, living, dining = _scene(client)
    hdr = _hdr(client)
    client.post(f"/api/items/{item['id']}/records",
                json={"qty": 1, "amount": 100, "room_id": living["id"]}, headers=hdr)
    client.post(f"/api/items/{item['id']}/records",
                json={"qty": 2, "amount": 200, "room_id": living["id"]}, headers=hdr)
    client.post(f"/api/items/{item['id']}/records",
                json={"qty": 4, "amount": 380, "room_id": dining["id"]}, headers=hdr)

    assert _cell(client, item["id"], living["id"]) == 3    # 1+2 正好买齐客厅
    assert _cell(client, item["id"], dining["id"]) == 4
    # 整条物料的状态照旧由总量与总实付推导
    assert client.get(f"/api/items/{item['id']}", headers=hdr).json()["status"] == "done"


def test_one_payment_can_cover_several_rooms(client):
    """一笔采购同时买几间的东西：按分配顺序依次抵扣，扣满一间再下一间。"""
    item, living, dining = _scene(client)
    hdr = _hdr(client)
    r = client.post(f"/api/items/{item['id']}/records",
                    json={"qty": 4, "amount": 400,
                          "room_ids": [living["id"], dining["id"]]}, headers=hdr)
    assert r.status_code == 200, r.text
    assert sorted(r.json()["records"][0]["room_ids"]) == sorted([living["id"], dining["id"]])

    assert _cell(client, item["id"], living["id"]) == 3    # 客厅 3 装满
    assert _cell(client, item["id"], dining["id"]) == 1    # 剩下的 1 个落在餐厅


def test_partial_payment_to_one_room(client):
    """某间只买了一部分：那一间不能算买齐。"""
    item, living, dining = _scene(client)
    client.post(f"/api/items/{item['id']}/records",
                json={"qty": 1, "amount": 100, "room_ids": [living["id"]]},
                headers=_hdr(client))
    assert _cell(client, item["id"], living["id"]) == 1    # 1/3
    assert _cell(client, item["id"], dining["id"]) == 0


def test_adding_quantity_later_unmarks_the_room(client):
    """买完之后又给那个分组加了数量：它应该变回"没买齐"。"""
    item, living, dining = _scene(client)
    hdr = _hdr(client)
    client.post(f"/api/items/{item['id']}/records",
                json={"qty": 3, "amount": 300, "room_ids": [living["id"]]}, headers=hdr)
    assert _cell(client, item["id"], living["id"]) == 3    # 3/3 买齐了

    # 给客厅加两个（改成 5 个）
    detail = client.get(f"/api/items/{item['id']}", headers=hdr).json()
    client.put(f"/api/items/{item['id']}", headers=hdr, json={
        "name": detail["name"], "unit": detail["unit"], "qty_total": 0,
        "price": detail["price"],
        "allocations": [{"room_id": living["id"], "qty": 5},
                        {"room_id": dining["id"], "qty": 4}],
        "records": [{"qty": 3, "amount": 300, "room_ids": [living["id"]]}],
        "base_rev": detail["rev"],
    })
    assert _cell(client, item["id"], living["id"]) == 3    # 覆盖还是那 3 个
    # 但 3 < 5，所以它不再是"已买齐"
    matrix = client.get("/api/matrix", headers=hdr).json()
    row = next(i for i in matrix["items"] if i["id"] == item["id"])
    cell = row["cells"][str(living["id"])]
    assert cell["paid_qty"] < cell["qty"]


def test_clearing_rooms_with_empty_list(client):
    item, living, dining = _scene(client)
    hdr = _hdr(client)
    rec = client.post(f"/api/items/{item['id']}/records",
                      json={"qty": 3, "amount": 300, "room_ids": [living["id"]]},
                      headers=hdr).json()["records"][0]
    assert rec["room_ids"] == [living["id"]]

    # 显式传空列表 = 清空涉及的分组
    cleared = client.put(f"/api/records/{rec['id']}", json={"room_ids": []}, headers=hdr)
    assert cleared.status_code == 200
    assert cleared.json()["records"][0]["room_ids"] == []
    assert _cell(client, item["id"], living["id"]) == 0    # 不再算在任何一間头上


def test_unassigned_payments_still_do_not_guess(client):
    """没写归属的多分组部分买齐：仍然一个都不标（老规矩不能丢）。"""
    item, living, dining = _scene(client)
    client.post(f"/api/items/{item['id']}/records",
                json={"qty": 5, "amount": 500}, headers=_hdr(client))
    assert _cell(client, item["id"], living["id"]) == 0
    assert _cell(client, item["id"], dining["id"]) == 0


def test_mixed_payments_use_what_is_known(client):
    """一部分写了归属、一部分没写：写了的照实算，没写的那部分不猜。"""
    item, living, dining = _scene(client)
    hdr = _hdr(client)
    client.post(f"/api/items/{item['id']}/records",
                json={"qty": 4, "amount": 400, "room_id": dining["id"]}, headers=hdr)
    client.post(f"/api/items/{item['id']}/records",
                json={"qty": 1, "amount": 100}, headers=hdr)   # 没写归属
    assert _cell(client, item["id"], dining["id"]) == 4
    assert _cell(client, item["id"], living["id"]) == 0


def test_single_group_item_still_spreads_unassigned(client):
    """只分给一个分组时，没写归属的付款照样按顺序抵（不存在先后问题）。"""
    room = client.post("/api/rooms", json={"name": "客厅", "sort": 0},
                       headers=_hdr(client)).json()
    item = client.post("/api/items", json={"name": "筒灯", "unit": "个", "price": 10},
                       headers=_hdr(client)).json()
    hdr = _hdr(client)
    client.put(f"/api/items/{item['id']}", headers=hdr, json={
        "name": "筒灯", "unit": "个", "qty_total": 0, "price": 10,
        "allocations": [{"room_id": room["id"], "qty": 4}], "records": [],
        "base_rev": client.get(f"/api/items/{item['id']}", headers=hdr).json()["rev"],
    })
    client.post(f"/api/items/{item['id']}/records", json={"qty": 1, "amount": 10},
                headers=hdr)
    assert _cell(client, item["id"], room["id"]) == 1


def test_room_must_belong_to_current_list(client):
    item, living, _ = _scene(client)
    other = client.post("/api/lists", json={"name": "年货清单"}).json()
    foreign_room = client.post("/api/rooms", json={"name": "别人家的分组", "sort": 0},
                              headers={"X-List-Id": str(other["id"])}).json()
    r = client.post(f"/api/items/{item['id']}/records",
                    json={"qty": 1, "amount": 100, "room_id": foreign_room["id"]},
                    headers=_hdr(client))
    assert r.status_code == 400 and "不属于当前清单" in r.json()["detail"]


def test_assignment_can_be_changed_or_cleared(client):
    item, living, dining = _scene(client)
    hdr = _hdr(client)
    rec = client.post(f"/api/items/{item['id']}/records",
                      json={"qty": 4, "amount": 400, "room_id": dining["id"]},
                      headers=hdr).json()["records"][0]

    # 改成客厅
    upd = client.put(f"/api/records/{rec['id']}",
                     json={"room_id": living["id"]}, headers=hdr)
    assert upd.status_code == 200
    assert upd.json()["records"][0]["room_id"] == living["id"]
    assert _cell(client, item["id"], living["id"]) == 3    # 客厅只有 3 个，抵 3
    assert _cell(client, item["id"], dining["id"]) == 0

    # 取消归属：显式传 null（不传这个字段才是"不改"）
    cleared = client.put(f"/api/records/{rec['id']}", json={"room_id": None}, headers=hdr)
    assert cleared.status_code == 200
    assert cleared.json()["records"][0]["room_id"] is None


def test_deleting_room_keeps_the_payment(client):
    """删分组时归属置空，付款记录本身不能丢 —— 那是钱。"""
    item, living, dining = _scene(client)
    hdr = _hdr(client)
    client.post(f"/api/items/{item['id']}/records",
                json={"qty": 4, "amount": 400, "room_id": dining["id"]}, headers=hdr)

    assert client.delete(f"/api/rooms/{dining['id']}", headers=hdr).status_code == 200
    got = client.get(f"/api/items/{item['id']}", headers=hdr).json()
    assert got["paid"] == 400                    # 钱还在
    assert got["records"][0]["room_id"] is None  # 归属没了


def test_excel_roundtrip_keeps_room_assignment(client):
    item, living, dining = _scene(client)
    hdr = _hdr(client)
    client.post(f"/api/items/{item['id']}/records",
                json={"qty": 4, "amount": 400, "date": "2026-09-14",
                      "room_id": dining["id"]}, headers=hdr)

    exported = client.get("/api/export", headers=hdr).content
    ws = openpyxl.load_workbook(io.BytesIO(exported))["采购记录"]
    header = [c.value for c in ws[1]]
    row = next(ws.iter_rows(min_row=2, values_only=True))
    assert dict(zip(header, row))["分组"] == "餐厅"

    r = client.post("/api/import",
                    files={"file": ("a.xlsx", io.BytesIO(exported),
                                    "application/vnd.openxmlformats-officedocument"
                                    ".spreadsheetml.sheet")},
                    data={"mode": "replace"}, headers=hdr)
    assert r.status_code == 200, r.text
    got = client.get(f"/api/items/{item['id']}", headers=hdr).json()
    saved = [r for r in got["records"] if r["qty"] == 4]
    assert saved and saved[0]["room_id"] == dining["id"]
    assert _cell(client, item["id"], dining["id"]) == 4


def test_amounts_are_untouched_by_assignment(client):
    """归属只是"这笔钱算在哪间"，不改变任何金额口径。"""
    item, living, dining = _scene(client)
    hdr = _hdr(client)
    client.post(f"/api/items/{item['id']}/records",
                json={"qty": 4, "amount": 400}, headers=hdr)
    before = client.get("/api/summary", headers=hdr).json()["totals"]

    rec = client.get(f"/api/items/{item['id']}", headers=hdr).json()["records"][0]
    client.put(f"/api/records/{rec['id']}", json={"room_id": dining["id"]}, headers=hdr)

    after = client.get("/api/summary", headers=hdr).json()["totals"]
    assert after == before
    session = SessionLocal()
    try:
        row = session.query(PurchaseRecord).one()
        assert compute.item_paid(session.query(Item).one()) == 400
        # 涉及的分组现在存在 record_rooms 里（可多选）
        assert [rr.room_id for rr in row.rooms] == [dining["id"]]
    finally:
        session.close()
