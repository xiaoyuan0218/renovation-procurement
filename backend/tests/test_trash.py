"""回收站：删掉的物料能捞回来，也能彻底删。

软删除的风险是"漏掉某个查询，被删的东西又冒出来"，所以这里挨个界面验一遍：
清单、矩阵、总览、导出、清单计数。
"""

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import auth
from app.db import SessionLocal, engine
from app.main import app
from app.models import (Allocation, Category, ExtraExpense, Item, ItemList,
                        PurchaseRecord, Room, User)
from app.seed import init_db

USER = "admin"
PASSWORD = "s3cret-pass"


@pytest.fixture()
def client():
    init_db()
    session = SessionLocal()
    for table in (ExtraExpense, Allocation, Item, PurchaseRecord, Room, Category,
                  User, ItemList):
        session.query(table).delete()
    session.commit()
    session.add(ItemList(name="采购清单", sort=0))
    session.commit()
    session.close()
    auth._failures.clear()
    with TestClient(app) as c:
        c.post("/api/auth/setup", json={"username": USER, "password": PASSWORD})
        yield c
    engine.dispose()


def _hdr(client):
    return {"X-List-Id": str(client.get("/api/lists").json()[0]["id"])}


def _mk(client, name="筒灯", qty=4, price=100):
    return client.post("/api/items", json={"name": name, "unit": "个",
                                           "qty_total": qty, "price": price},
                       headers=_hdr(client)).json()


def _with_allocation(client, name="筒灯"):
    """建一条带分组分配 + 一笔付款的物料，用来验证"删了别把历史丢了"。"""
    item = _mk(client, name)
    room = client.post("/api/rooms", json={"name": f"{name}的房间", "sort": 0},
                       headers=_hdr(client)).json()
    client.put("/api/matrix/cell",
               json={"item_id": item["id"], "room_id": room["id"], "qty": 3},
               headers=_hdr(client))
    client.post(f"/api/items/{item['id']}/records",
                json={"qty": 1, "amount": 80, "date": "2026-09-14"},
                headers=_hdr(client))
    return item


def test_deleted_item_disappears_from_every_view(client):
    item = _with_allocation(client)
    other = _mk(client, "网线")
    totals_before = client.get("/api/summary", headers=_hdr(client)).json()["totals"]

    r = client.delete(f"/api/items/{item['id']}", headers=_hdr(client))
    assert r.status_code == 200 and r.json()["trashed"] is True

    # 清单里没了、单条访问 404、矩阵里没了
    assert [i["name"] for i in client.get("/api/items", headers=_hdr(client)).json()] == ["网线"]
    assert client.get(f"/api/items/{item['id']}", headers=_hdr(client)).status_code == 404
    matrix = client.get("/api/matrix", headers=_hdr(client)).json()
    assert [i["name"] for i in matrix["items"]] == ["网线"]
    # 总览的金额跟着少掉这一条（它原来占 300）
    totals_after = client.get("/api/summary", headers=_hdr(client)).json()["totals"]
    assert totals_after["list_total"] == totals_before["list_total"] - 300
    assert totals_after["item_count"] == 1
    # 清单计数里也不算它
    assert client.get("/api/lists").json()[0]["item_count"] == 1
    # 导出里没有它
    wb = openpyxl.load_workbook(io.BytesIO(client.get("/api/export", headers=_hdr(client)).content))
    names = [row[1] for row in wb["物料汇总"].iter_rows(min_row=2, values_only=True)]
    assert names == ["网线"]
    assert other["id"]  # 另一条不受影响


def test_batch_delete_is_soft_too(client):
    a = _mk(client, "A")
    b = _mk(client, "B")
    r = client.post("/api/items/batch/delete", json={"ids": [a["id"], b["id"]]},
                    headers=_hdr(client))
    assert r.json()["deleted"] == 2 and r.json()["trashed"] is True
    assert client.get("/api/items", headers=_hdr(client)).json() == []
    assert len(client.get("/api/trash", headers=_hdr(client)).json()) == 2


def test_trash_keeps_allocations_and_records(client):
    item = _with_allocation(client)
    client.delete(f"/api/items/{item['id']}", headers=_hdr(client))

    rows = client.get("/api/trash", headers=_hdr(client)).json()
    assert len(rows) == 1
    row = rows[0]
    assert row["name"] == "筒灯"
    assert row["deleted_at"]           # 有删除时间
    assert len(row["allocations"]) == 1
    assert len(row["records"]) == 1
    assert row["paid"] == 80           # 付款历史还在


def test_restore_brings_it_back_intact(client):
    item = _with_allocation(client)
    before = client.get(f"/api/items/{item['id']}", headers=_hdr(client)).json()
    client.delete(f"/api/items/{item['id']}", headers=_hdr(client))

    r = client.post(f"/api/trash/{item['id']}/restore", headers=_hdr(client))
    assert r.status_code == 200, r.text
    after = client.get(f"/api/items/{item['id']}", headers=_hdr(client)).json()
    # 除了 rev（恢复会顶高它），其余一字不差
    assert {k: v for k, v in after.items() if k != "rev"} == \
           {k: v for k, v in before.items() if k != "rev"}
    assert after["rev"] > before["rev"]
    assert client.get("/api/trash", headers=_hdr(client)).json() == []


def test_purge_removes_for_good(client):
    item = _with_allocation(client)
    client.delete(f"/api/items/{item['id']}", headers=_hdr(client))

    assert client.delete(f"/api/trash/{item['id']}", headers=_hdr(client)).status_code == 200
    session = SessionLocal()
    try:
        assert session.query(Item).count() == 0
        assert session.query(Allocation).count() == 0
        assert session.query(PurchaseRecord).count() == 0
    finally:
        session.close()


def test_purge_all_clears_the_trash(client):
    for name in ("A", "B"):
        _mk(client, name)
    ids = [i["id"] for i in client.get("/api/items", headers=_hdr(client)).json()]
    client.post("/api/items/batch/delete", json={"ids": ids}, headers=_hdr(client))
    assert len(client.get("/api/trash", headers=_hdr(client)).json()) == 2

    assert client.delete("/api/trash", headers=_hdr(client)).json()["deleted"] == 2
    assert client.get("/api/trash", headers=_hdr(client)).json() == []


def test_trashed_item_cannot_be_edited_directly(client):
    item = _mk(client)
    client.delete(f"/api/items/{item['id']}", headers=_hdr(client))
    hdr = _hdr(client)
    assert client.get(f"/api/items/{item['id']}", headers=hdr).status_code == 404
    assert client.patch(f"/api/items/{item['id']}", json={"note": "x"},
                        headers=hdr).status_code == 404
    assert client.put(f"/api/items/{item['id']}",
                      json={"name": "改过", "unit": "个", "qty_total": 1, "price": 1},
                      headers=hdr).status_code == 404
    assert client.post(f"/api/items/{item['id']}/records",
                       json={"qty": 1, "amount": 1}, headers=hdr).status_code == 404
    # 恢复之后又能编辑了
    client.post(f"/api/trash/{item['id']}/restore", headers=hdr)
    assert client.patch(f"/api/items/{item['id']}", json={"note": "回来了"},
                        headers=hdr).status_code == 200


def test_trash_is_isolated_between_lists(client):
    first = client.get("/api/lists").json()[0]
    other = client.post("/api/lists", json={"name": "年货清单"}).json()
    hdr_other = {"X-List-Id": str(other["id"])}
    item = _mk(client)
    client.delete(f"/api/items/{item['id']}", headers=_hdr(client))

    assert client.get("/api/trash", headers=hdr_other).json() == []
    assert client.post(f"/api/trash/{item['id']}/restore",
                       headers=hdr_other).status_code == 404
    assert client.delete(f"/api/trash/{item['id']}",
                         headers=hdr_other).status_code == 404
    assert len(client.get("/api/trash", headers=_hdr(client)).json()) == 1


def test_deleting_list_takes_the_trash_with_it(client):
    first = client.get("/api/lists").json()[0]
    other = client.post("/api/lists", json={"name": "年货清单"}).json()
    item = _mk(client)
    client.delete(f"/api/items/{item['id']}", headers=_hdr(client))

    assert client.delete(f"/api/lists/{first['id']}").status_code == 200
    session = SessionLocal()
    try:
        assert session.query(Item).count() == 0   # 回收站里的也跟着走了
    finally:
        session.close()


def test_import_replace_leaves_the_trash_alone(client):
    """覆盖导入只替换"看得见的条目"，不动回收站 —— 用户还没决定要不要恢复。"""
    item = _mk(client, "旧物料")
    client.delete(f"/api/items/{item['id']}", headers=_hdr(client))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "物料汇总"
    ws.append(["类目", "物料名称", "单位", "数量", "单价", "优惠单价",
               "日常价", "实付数量", "实付金额", "未付数量", "未付金额",
               "已购", "备注"])
    ws.append(["照明", "新物料", "个", 2, 99, None, None, None, None, None,
               None, "否", ""])
    buf = io.BytesIO()
    wb.save(buf)

    r = client.post("/api/import",
                    files={"file": ("x.xlsx", buf.getvalue(),
                                    "application/vnd.openxmlformats-officedocument"
                                    ".spreadsheetml.sheet")},
                    data={"mode": "replace"}, headers=_hdr(client))
    assert r.status_code == 200, r.text
    assert [i["name"] for i in client.get("/api/items", headers=_hdr(client)).json()] == ["新物料"]
    assert [i["name"] for i in client.get("/api/trash", headers=_hdr(client)).json()] == ["旧物料"]
