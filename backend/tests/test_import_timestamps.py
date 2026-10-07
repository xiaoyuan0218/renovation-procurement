"""导入导出要带着「添加时间 / 修改时间」走。

为什么值得单独盯一组测试：这几列是同步判"谁改得更近"的依据。导出一份、再灌
回去（搬家、备份、手机电脑互传都这么用）如果时间被刷成"刚刚"，同步就会把所有
行都当成新改动，较旧的那份反而盖掉对方较新的改动。另外「物料ID」挪到了第一列，
老文件（ID 在最后一列）还得照样能读。
"""

import io
from datetime import datetime

import openpyxl
import pytest

from app.db import SessionLocal, engine
from app.models import (Allocation, Category, ExtraExpense, Item, ItemList,
                        PurchaseRecord, RecordRoom, Room)
from app.seed import init_db
from app.services import codes, excel_io


@pytest.fixture()
def db():
    init_db()
    session = SessionLocal()
    for table in (RecordRoom, Allocation, PurchaseRecord, ExtraExpense, Item,
                  Room, Category):
        session.query(table).delete()
    session.commit()
    yield session
    session.close()
    engine.dispose()


@pytest.fixture()
def seeded(db):
    """一条物料 + 一条布点 + 一笔采购 + 一笔额外费用，时间都设成固定值。"""
    lst = db.query(ItemList).order_by(ItemList.sort, ItemList.id).first()
    room = Room(name="客厅", sort=0, list_id=lst.id)
    cat = Category(name="照明", sort=0, list_id=lst.id)
    db.add_all([room, cat])
    db.flush()

    made = datetime(2026, 3, 4, 5, 6, 7)
    changed = datetime(2026, 4, 5, 6, 7, 8)
    item = Item(name="筒灯", brand="松下", unit="个", qty_total=6, price=30,
                category_id=cat.id, list_id=lst.id,
                created_at=made, updated_at=changed)
    db.add(item)
    db.flush()
    db.add(Allocation(item_id=item.id, room_id=room.id, qty=6))
    db.add(PurchaseRecord(item_id=item.id, qty=6, amount=180.5, date="2026-03-10",
                          vendor="京东", created_at=made, updated_at=changed))
    db.add(ExtraExpense(list_id=lst.id, kind="运费", amount=120, date="2026-03-10",
                        created_at=made, updated_at=changed))
    db.commit()
    return {"list_id": lst.id, "made": made, "changed": changed}


def _headers(data: bytes, sheet: str):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    return [c.value for c in wb[sheet][1]]


def _cells(data: bytes, sheet: str):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb[sheet]
    header = [c.value for c in ws[1]]
    return header, [dict(zip(header, [c.value for c in row]))
                    for row in ws.iter_rows(min_row=2)]


def test_item_id_is_the_first_column(db, seeded):
    """「物料ID」在第一列：它才是物料的身份，用户对账时一眼能找到它。"""
    data = excel_io.export_xlsx(db, seeded["list_id"])
    assert _headers(data, "物料汇总")[0] == "物料ID"
    assert _headers(data, "布点明细")[0] == "物料ID"
    assert _headers(data, "采购记录")[0] == "物料ID"


def test_template_keeps_the_same_columns(db, seeded):
    """模板与导出逐列一致（共用同一份表头常量，这里守住别被改回去）。"""
    tpl = _headers(excel_io.build_template(), "物料汇总")
    exp = _headers(excel_io.export_xlsx(db, seeded["list_id"]), "物料汇总")
    assert tpl == exp


def test_export_carries_timestamps(db, seeded):
    """导出要带上两个时间，格式与接口一致（UTC、YYYY-MM-DD HH:MM:SS）。"""
    data = excel_io.export_xlsx(db, seeded["list_id"])
    _, items = _cells(data, "物料汇总")
    assert items[0]["添加时间"] == "2026-03-04 05:06:07"
    assert items[0]["修改时间"] == "2026-04-05 06:07:08"

    _, records = _cells(data, "采购记录")
    assert records[0]["添加时间"] == "2026-03-04 05:06:07"

    _, expenses = _cells(data, "额外费用")
    assert expenses[0]["添加时间"] == "2026-03-04 05:06:07"


def test_export_reimport_keeps_timestamps(db, seeded):
    """导出再灌回去，时间一模一样 —— 这是"回灌零变化"的一部分。"""
    data = excel_io.export_xlsx(db, seeded["list_id"])
    excel_io.import_template(db, data, mode="replace", list_id=seeded["list_id"])
    db.expire_all()

    item = db.query(Item).filter(Item.list_id == seeded["list_id"]).one()
    assert item.created_at == seeded["made"]
    assert item.updated_at == seeded["changed"]

    record = db.query(PurchaseRecord).one()
    assert record.created_at == seeded["made"]
    assert record.updated_at == seeded["changed"]

    expense = db.query(ExtraExpense).one()
    assert expense.created_at == seeded["made"]
    assert expense.updated_at == seeded["changed"]


def test_mobile_format_timestamp_with_microseconds(db, seeded):
    """手机导出的时间带微秒，电脑端也要认得（否则回灌会把时间丢掉）。"""
    data = excel_io.export_xlsx(db, seeded["list_id"])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["物料汇总"]
    header = [c.value for c in ws[1]]
    ws.cell(row=2, column=header.index("添加时间") + 1,
            value="2026-03-04 05:06:07.123456")
    buf = io.BytesIO()
    wb.save(buf)

    excel_io.import_template(db, buf.getvalue(), mode="replace",
                             list_id=seeded["list_id"])
    db.expire_all()
    item = db.query(Item).filter(Item.list_id == seeded["list_id"]).one()
    assert item.created_at == datetime(2026, 3, 4, 5, 6, 7, 123456)


def test_old_file_without_timestamp_columns_still_imports(db, seeded):
    """老文件没有这两列：照常导入，时间交给数据库默认值，不报错。"""
    data = excel_io.export_xlsx(db, seeded["list_id"])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["物料汇总"]
    header = [c.value for c in ws[1]]
    for name in ("添加时间", "修改时间"):
        ws.delete_cols(header.index(name) + 1)
        header = [c.value for c in ws[1]]
    buf = io.BytesIO()
    wb.save(buf)

    report = excel_io.import_template(db, buf.getvalue(), mode="replace",
                                      list_id=seeded["list_id"])
    assert report.get("warnings") == [], report.get("warnings")
    db.expire_all()
    item = db.query(Item).filter(Item.list_id == seeded["list_id"]).one()
    assert item.created_at is not None   # 数据库默认值兜底，不是空


def test_garbage_timestamp_is_ignored(db, seeded):
    """时间列填了认不出来的东西：忽略它，不报错也不写一个假时间。"""
    data = excel_io.export_xlsx(db, seeded["list_id"])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["物料汇总"]
    header = [c.value for c in ws[1]]
    ws.cell(row=2, column=header.index("修改时间") + 1, value="昨天下午")
    buf = io.BytesIO()
    wb.save(buf)

    report = excel_io.import_template(db, buf.getvalue(), mode="replace",
                                      list_id=seeded["list_id"])
    assert report.get("warnings") == [], report.get("warnings")
    db.expire_all()
    item = db.query(Item).filter(Item.list_id == seeded["list_id"]).one()
    assert item.updated_at is not None   # 走了数据库默认值，而不是崩掉
