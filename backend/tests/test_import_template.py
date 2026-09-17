import io

import openpyxl
import pytest

from app.db import SessionLocal, engine
from app.models import Allocation, Category, Item, PurchaseRecord, Room
from app.seed import init_db
from app.services import excel_io


@pytest.fixture()
def db():
    init_db()
    session = SessionLocal()
    for table in (Allocation, Item, PurchaseRecord, Room, Category):
        session.query(table).delete()
    session.commit()
    yield session
    session.close()
    engine.dispose()


def _flat_wb():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "物料汇总"
    ws.append(["类目", "物料名称", "单位", "数量", "单价", "日常单价",
               "日常价", "实付数量", "实付金额", "未付数量", "未付金额",
               "已购", "备注"])
    ws.append(["照明", "筒灯", "个", 10, 49, 39.39, None, None, None, None, None, None, "否", ""])
    ws.append(["网络", "网线", "米", 150, 4, None, None, None, None, None, None, None, "否", ""])
    ws2 = wb.create_sheet("布点明细")
    ws2.append(["物料名称", "房间", "数量", "单价", "备注"])
    ws2.append(["筒灯", "客厅", 3, None, ""])
    ws2.append(["筒灯", "餐厅", 4, 59, "餐厅用贵的"])
    ws2.append(["网线", "客厅", 0, None, ""])  # 0 数量应被忽略
    ws3 = wb.create_sheet("采购记录")
    ws3.append(["物料名称", "实付数量", "实付金额", "付款日期", "备注"])
    ws3.append(["筒灯", 3, 150, "2026-09-01", "第一批"])
    ws3.append(["筒灯", 4, 236.09, "2026-09-10", "第二批"])
    ws3.append(["网线", 150, 661.2, None, ""])
    ws3.append(["网线", 0, 0, None, ""])  # 空记录应被忽略
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_template_import_replace(db):
    report = excel_io.import_template(db, _flat_wb(), mode="replace")
    assert report["items_created"] == 2
    assert report["allocations"] == 2  # 0 数量行被忽略
    assert report["records"] == 3  # 空记录被忽略

    db.expire_all()  # 清掉 session 内已加载的关系缓存
    items = {i.name: i for i in db.query(Item).all()}
    deng = items["筒灯"]
    # 有布点 → 总量以布点合计为准 (3+4)
    assert deng.allocations and len(deng.allocations) == 2
    from app.services import compute
    d = compute.item_dict(deng)
    assert d["total_qty"] == 7
    # 布点单价覆盖：3×49 + 4×59
    assert d["list_total"] == 3 * 49 + 4 * 59
    assert d["discount_total"] == pytest.approx(39.39 * 7, abs=0.001)
    # 两笔采购记录：数量金额分别累加，实付单价=金额÷数量
    assert d["paid_qty"] == 7
    assert d["paid"] == pytest.approx(150 + 236.09, abs=0.01)
    assert d["paid_price"] == pytest.approx((150 + 236.09) / 7, abs=0.01)
    assert d["status"] == "done"
    assert len(d["records"]) == 2 and d["records"][0]["note"] == "第一批"

    wire = items["网线"]
    dw = compute.item_dict(wire)
    assert wire.qty_total == 150
    assert wire.discount_price is None  # 留空 → 按原价
    assert dw["paid_qty"] == 150 and dw["status"] == "done"
    assert compute.item_discount_total(wire) == 600
    assert compute.item_unpaid(wire) == 0  # 网线已记录一笔 661.2 → 数量150全覆盖


def test_template_import_merge(db):
    excel_io.import_template(db, _flat_wb(), mode="replace")
    # 再导一次 merge：同名物料更新而不是新建
    report = excel_io.import_template(db, _flat_wb(), mode="merge")
    assert report["items_matched"] == 2
    assert report["items_created"] == 0
    assert db.query(Item).count() == 2
    assert db.query(Allocation).count() == 2


def test_template_import_bad_format(db):
    wb = openpyxl.Workbook()
    wb.active.append(["随便什么表"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError):
        excel_io.import_template(db, buf.getvalue(), mode="replace")


def test_export_roundtrip(db):
    excel_io.import_template(db, _flat_wb(), mode="replace")
    data = excel_io.export_xlsx(db)
    # 导出的文件能按模板回灌
    report = excel_io.import_template(db, data, mode="replace")
    assert report["items_created"] == 2
    assert report["allocations"] == 2
    from app.services import compute
    deng = db.query(Item).filter(Item.name == "筒灯").one()
    d = compute.item_dict(deng)
    assert d["list_total"] == 3 * 49 + 4 * 59


def test_export_has_split_columns(db):
    """导出带上三段拆分的三列，位置在「未付金额」和「已购」之间，数值与 compute 一致。"""
    import openpyxl

    from app.services import compute
    excel_io.import_template(db, _flat_wb(), mode="replace")
    data = excel_io.export_xlsx(db)
    ws = openpyxl.load_workbook(io.BytesIO(data))["物料汇总"]

    header = [c.value for c in ws[1]]
    assert header[header.index("未付金额") + 1: header.index("未付金额") + 4] == \
        ["日常价未付", "实际优惠", "日常价优惠"]
    assert header.index("已购") == header.index("日常价优惠") + 1

    # 逐行核对：导出的三列就是 compute 的结果，且等式在表格里成立
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(header, row))
        item = db.query(Item).filter(Item.name == rec["物料名称"]).one()
        d = compute.item_dict(item)
        assert rec["日常价未付"] == d["daily_unpaid"]
        assert rec["实际优惠"] == d["actual_discount"]
        assert rec["日常价优惠"] == d["daily_discount"]


def test_export_split_identity_holds_per_row(db):
    """有房间覆盖价时也要成立：已付 + 实际优惠 + 未付 = 原价合计（= 行内各房间覆盖价之和）。"""
    import openpyxl

    from app.services import compute
    excel_io.import_template(db, _flat_wb(), mode="replace")
    ws = openpyxl.load_workbook(io.BytesIO(excel_io.export_xlsx(db)))["物料汇总"]
    header = [c.value for c in ws[1]]

    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(header, row))
        item = db.query(Item).filter(Item.name == rec["物料名称"]).one()
        d = compute.item_dict(item)
        assert round(d["paid"] + d["actual_discount"] + d["unpaid"], 2) == d["list_total"]
        assert round(d["paid"] + d["daily_discount"] + d["daily_unpaid"], 2) == d["discount_total"]


def test_template_keeps_same_header_as_export(db):
    """模板必须与导出的表头一致，否则「导出当模板用」这条使用方式会断。"""
    import openpyxl
    tpl = openpyxl.load_workbook(io.BytesIO(excel_io.build_template()))["物料汇总"]
    exp = openpyxl.load_workbook(io.BytesIO(excel_io.export_xlsx(db)))["物料汇总"]
    assert [c.value for c in tpl[1]] == [c.value for c in exp[1]]
    # 示例行的列数也要对齐，否则用户不删示例行就会导入串列的数据
    for row in tpl.iter_rows(min_row=2, values_only=True):
        assert len(row) == len([c.value for c in tpl[1]])


def _full_snapshot(db):
    """所有会进金额口径的量 + 分配/记录明细，用于判断一次导入有没有改变任何东西。"""
    from app.services import compute
    db.expire_all()
    views = [compute.item_dict(i) for i in db.query(Item).all()]
    return {
        "items": db.query(Item).count(),
        "rooms": db.query(Room).count(),
        "cats": db.query(Category).count(),
        "list_total": round(sum(v["list_total"] for v in views), 2),
        "discount_total": round(sum(v["discount_total"] for v in views), 2),
        "paid": round(sum(v["paid"] for v in views), 2),
        "unpaid": round(sum(v["unpaid"] for v in views), 2),
        "per_item": {v["id"]: (v["name"], v["total_qty"], v["list_total"], v["paid"], v["unpaid"])
                     for v in views},
        # 明细也要逐字段一致：覆盖价是不是 None 也算差异（留空=跟随物料单价，
        # 被写成具体数字后，以后改物料单价这一行就不跟着动了）
        "allocs": sorted((a.item_id, a.room_id, a.qty, a.price_override, a.note or "")
                         for a in db.query(Allocation).all()),
        "records": sorted((r.item_id, r.qty, r.amount, r.date or "", r.note or "",
                           r.vendor or "", r.order_no or "")
                          for r in db.query(PurchaseRecord).all()),
    }


def _build_duplicate_name_db(db):
    """两条同名物料 + 各自的分配与记录 —— 导出回灌必须原样往返。

    这是真实数据里踩到的坑：导入按名称匹配，同名物料会让第二个物料的分配
    不被清空、而两份分配又都挂到第一个物料上，导致分配变多、金额上涨。
    第一条留一个「无覆盖价」的分配，用来盯住覆盖价在往返中被写实的问题。
    """
    room_a = Room(name="客厅", sort=0)
    room_b = Room(name="餐厅", sort=1)
    db.add_all([room_a, room_b])
    db.flush()
    cat = Category(name="照明", sort=0)
    db.add(cat)
    db.flush()

    first = Item(name="易来灯带控制器", unit="个", category_id=cat.id,
                 price=99, discount_price=80)
    db.add(first)
    db.flush()
    db.add(Allocation(item_id=first.id, room_id=room_a.id, qty=3, price_override=119))
    db.add(Allocation(item_id=first.id, room_id=room_b.id, qty=3))          # 无覆盖价
    first.records.append(PurchaseRecord(qty=1, amount=200, note="定金"))
    first.records.append(PurchaseRecord(qty=1, amount=180, date="2026-09-10"))

    second = Item(name="易来灯带控制器", unit="个", category_id=cat.id,
                  price=99, discount_price=80)
    db.add(second)
    db.flush()
    db.add(Allocation(item_id=second.id, room_id=room_a.id, qty=1))
    second.records.append(PurchaseRecord(qty=1, amount=50))
    db.commit()


def test_export_reimport_is_noop_with_duplicate_names(db):
    """导出 → 合并回灌必须是零变化，同名物料也不例外。"""
    _build_duplicate_name_db(db)
    before = _full_snapshot(db)

    data = excel_io.export_xlsx(db)
    excel_io.import_template(db, data, mode="merge")

    after = _full_snapshot(db)
    assert after == before, {k: (before[k], after[k]) for k in before if before[k] != after[k]}


def test_export_reimport_is_noop_twice(db):
    """连回灌两次也不能累积（幂等，不只往返一次）。"""
    _build_duplicate_name_db(db)
    before = _full_snapshot(db)
    for _ in range(2):
        data = excel_io.export_xlsx(db)
        excel_io.import_template(db, data, mode="merge")
    assert _full_snapshot(db) == before


def test_allocation_without_override_stays_without_override(db):
    """「跟随物料单价」的分配，往返后仍然是跟随，不会被写成一个固定价。"""
    _build_duplicate_name_db(db)
    data = excel_io.export_xlsx(db)
    excel_io.import_template(db, data, mode="merge")
    db.expire_all()

    room_b = db.query(Room).filter(Room.name == "餐厅").one()
    inherited = db.query(Allocation).filter(Allocation.room_id == room_b.id).one()
    assert inherited.price_override is None

    # 把物料单价改掉，跟随的那行应声而变，固定价那行不动
    first = db.query(Item).filter(Item.name == "易来灯带控制器").order_by(Item.id).first()
    room_a = db.query(Room).filter(Room.name == "客厅").one()
    overridden = (db.query(Allocation)
                  .filter(Allocation.item_id == first.id, Allocation.room_id == room_a.id).one())
    first.price = 200
    db.commit()
    db.expire_all()
    assert overridden.price_override == 119


def test_export_carries_item_id(db):
    """三个 sheet 都要带上物料ID列。"""
    _build_duplicate_name_db(db)
    wb = openpyxl.load_workbook(io.BytesIO(excel_io.export_xlsx(db)))
    for sheet in ("物料汇总", "布点明细", "采购记录"):
        header = [c.value for c in wb[sheet][1]]
        assert "物料ID" in header, f"{sheet} 缺少物料ID列"
        ids = [row[-1] for row in wb[sheet].iter_rows(min_row=2, values_only=True)]
        assert all(ids), f"{sheet} 有行没写物料ID"


def test_legacy_file_without_item_id_still_imports(db):
    """老文件（没有物料ID列）仍能导入，且名称唯一时结果同样精确。"""
    room = Room(name="客厅", sort=0)
    db.add(room)
    db.flush()
    item = Item(name="筒灯", price=99, discount_price=79.5, unit="个", qty_total=6)
    db.add(item)
    db.flush()
    db.add(Allocation(item_id=item.id, room_id=room.id, qty=6, price_override=119))
    item.records.append(PurchaseRecord(qty=1, amount=200))
    db.commit()

    data = excel_io.export_xlsx(db)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    for sheet in ("物料汇总", "布点明细", "采购记录"):
        ws = wb[sheet]
        ws.delete_cols([c.value for c in ws[1]].index("物料ID") + 1)
    buf = io.BytesIO()
    wb.save(buf)

    before = _full_snapshot(db)
    excel_io.import_template(db, buf.getvalue(), mode="merge")
    assert _full_snapshot(db) == before
